import io
import re
import tempfile
import zipfile
from collections import Counter
from datetime import datetime, timezone
from secrets import token_urlsafe
from typing import List

from pathlib import Path

import qrcode
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .config import get_settings
from .database import SessionLocal, check_database, get_db, init_database
from .models import AttendanceLog, Event, Participant, Table, User, Vendor
from .security import create_token, decode_token, hash_password, verify_password

settings = get_settings()
app = FastAPI(
    title="Gatherly — Event Registration & Check-in API",
    version="2.0.0",
    docs_url=None if settings.is_production else "/docs",
    redoc_url=None if settings.is_production else "/redoc",
    openapi_url=None if settings.is_production else "/openapi.json",
)
_origins = settings.cors_origin_list
_is_wildcard = _origins == ["*"]
app.add_middleware(CORSMiddleware, allow_origins=_origins, allow_credentials=not _is_wildcard, allow_methods=["*"], allow_headers=["*"])
bearer = HTTPBearer(auto_error=False)

# ─── Pydantic Schemas ──────────────────────────────────────────────────────────

class LoginInput(BaseModel):
    email: str
    password: str

class EventInput(BaseModel):
    name: str = Field(min_length=2, max_length=180)
    description: str = ""
    event_date: str
    start_time: str = "08:00"
    end_time: str = "17:00"
    location: str = ""
    status: str = "active"

class VendorInput(BaseModel):
    company_name: str = Field(min_length=2, max_length=180)
    category: str = "General"
    contact_name: str = ""
    phone: str = ""
    email: str = ""

class ParticipantInput(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    vendor_id: int | None = None
    table_id: int | None = None
    seat_number: str | None = None
    position: str = ""
    phone: str = ""
    email: str = ""

class GroupParticipantInput(BaseModel):
    names: list[str] = Field(min_length=2)
    vendor_id: int | None = None
    table_id: int | None = None
    seat_number: str | None = None
    position: str = ""
    phone: str = ""
    email: str = ""

class GroupUpdateInput(BaseModel):
    names: list[str] = Field(min_length=2)
    vendor_id: int | None = None
    table_id: int | None = None
    seat_number: str | None = None
    position: str = ""
    phone: str = ""
    email: str = ""

class ScanInput(BaseModel):
    token: str = Field(min_length=4)
    notes: str = ""
    attended_by: str | None = None
    is_substitute: bool = False
    participant_ids: list[int] | None = None
    attended_by_by_participant: dict[int, str] | None = None

class ManualCheckinInput(BaseModel):
    participant_id: int
    notes: str = ""
    attended_by: str | None = None
    is_substitute: bool = False
    participant_ids: list[int] | None = None
    attended_by_by_participant: dict[int, str] | None = None

class PreviewScanInput(BaseModel):
    token: str | None = None
    participant_id: int | None = None

class UndoCheckinInput(BaseModel):
    participant_id: int
    notes: str = ""
    undo_delegasi: bool = False

class TableInput(BaseModel):
    table_number: str = Field(min_length=1, max_length=20)
    table_label: str = ""
    capacity: int = 8
    zone: str = "Regular"

class TableBulkInput(BaseModel):
    prefix: str = Field(min_length=1, max_length=10)
    count: int = Field(ge=1, le=100)
    capacity: int = 8
    zone: str = "Regular"

class SeatAssignInput(BaseModel):
    table_id: int | None = None
    seat_number: str | None = None

class ImportConfirmInput(BaseModel):
    rows: list[dict]
    duplicate_action: str = "skip"

class UserInput(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(min_length=5, max_length=255)
    password: str = Field(min_length=6, max_length=100)
    role: str = "operator"

class UserUpdateInput(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(min_length=5, max_length=255)
    role: str = "operator"
    is_active: bool = True

class PasswordChangeInput(BaseModel):
    current_password: str
    new_password: str = Field(min_length=6, max_length=100)

class EventCopyInput(BaseModel):
    name: str = Field(min_length=2, max_length=180)
    event_date: str
    copy_vendors: bool = True
    copy_tables: bool = True
    copy_participants: bool = False
    status: str = "active"

class UserOut(BaseModel):
    model_config = ConfigDict(from_attributes=True)
    id: int
    name: str
    email: str
    role: str
    is_active: bool = True

# ─── Helpers ────────────────────────────────────────────────────────────────────

def serialize_event(event: Event, db: Session) -> dict:
    total = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event.id)) or 0
    checked = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event.id, Participant.check_in_at.is_not(None))) or 0
    return {"id": event.id, "name": event.name, "description": event.description, "event_date": event.event_date, "start_time": event.start_time, "end_time": event.end_time, "location": event.location, "status": event.status, "total_participants": total, "checked_in": checked}


def group_members(item: Participant, db: Session) -> list[Participant]:
    if not item.qr_group_token:
        return [item]
    return list(db.scalars(
        select(Participant).where(
            Participant.event_id == item.event_id,
            Participant.qr_group_token == item.qr_group_token,
        ).order_by(Participant.id)
    ).all())


def effective_qr_token(item: Participant) -> str:
    return item.qr_group_token or item.qr_token


def build_group_lookup(participants: list[Participant]) -> dict[str, list[Participant]]:
    lookup: dict[str, list[Participant]] = {}
    for p in participants:
        if not p.qr_group_token:
            continue
        lookup.setdefault(p.qr_group_token, []).append(p)
    for members in lookup.values():
        members.sort(key=lambda m: m.id)
    return lookup


def serialize_participant(item: Participant, db: Session, group_lookup: dict[str, list[Participant]] | None = None) -> dict:
    vendor = db.get(Vendor, item.vendor_id) if item.vendor_id else None
    table = db.get(Table, item.table_id) if item.table_id else None
    if item.qr_group_token and group_lookup is not None:
        members = group_lookup.get(item.qr_group_token, [item])
    elif item.qr_group_token:
        members = group_members(item, db)
    else:
        members = [item]
    is_group = bool(item.qr_group_token) and len(members) > 1
    return {
        "id": item.id, "name": item.name, "position": item.position,
        "phone": item.phone, "email": item.email, "qr_token": effective_qr_token(item),
        "qr_group_token": item.qr_group_token,
        "is_group": is_group,
        "group_size": len(members) if is_group else 1,
        "group_names": [m.name for m in members] if is_group else [item.name],
        "attendance_status": item.attendance_status,
        "check_in_at": item.check_in_at, "check_in_method": item.check_in_method,
        "check_out_at": item.check_out_at,
        "attended_by": item.attended_by,
        "is_substitute": bool(item.attended_by),
        "vendor_id": item.vendor_id, "company_name": vendor.company_name if vendor else "Unassigned",
        "table_id": item.table_id,
        "table_number": table.table_number if table else None,
        "table_zone": table.zone if table else None,
        "seat_number": item.seat_number,
    }


def current_user(credentials: HTTPAuthorizationCredentials = Depends(bearer), db: Session = Depends(get_db)) -> User:
    if not credentials:
        raise HTTPException(status_code=401, detail="Authentication required")
    try:
        payload = decode_token(credentials.credentials)
        user = db.get(User, int(payload["sub"]))
    except Exception as exc:
        raise HTTPException(status_code=401, detail="Invalid or expired token") from exc
    if not user or not user.is_active:
        raise HTTPException(status_code=401, detail="User is inactive")
    return user


ROLE_HIERARCHY = {"superadmin": 4, "admin": 3, "operator": 2, "viewer": 1}

def require_role(minimum_role: str):
    def checker(user: User = Depends(current_user)):
        user_level = ROLE_HIERARCHY.get(user.role, 0)
        required_level = ROLE_HIERARCHY.get(minimum_role, 0)
        if user_level < required_level:
            raise HTTPException(status_code=403, detail="Akses ditolak. Role tidak mencukupi.")
        return user
    return checker


# QR slot inside template-card-qr.png (inner white area of the blue frame)
_QR_TEMPLATE_PATH = Path(__file__).resolve().parent / "assets" / "template-card-qr.png"
_QR_SLOT = (126, 329, 517, 708)  # left, top, right, bottom
_QR_FILL = (6, 75, 152)  # Mandiri Taspen blue from template
_QR_PADDING = 28


def generate_qr_image(token: str) -> io.BytesIO:
    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=12, border=1)
    qr.add_data(token)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color=_QR_FILL, back_color="white").convert("RGBA")

    if not _QR_TEMPLATE_PATH.exists():
        buf = io.BytesIO()
        qr_img.save(buf, format="PNG")
        buf.seek(0)
        return buf

    card = Image.open(_QR_TEMPLATE_PATH).convert("RGBA")
    left, top, right, bottom = _QR_SLOT
    slot_w = right - left - (_QR_PADDING * 2)
    slot_h = bottom - top - (_QR_PADDING * 2)
    qr_size = min(slot_w, slot_h)
    qr_img = qr_img.resize((qr_size, qr_size), Image.Resampling.NEAREST)

    offset_x = left + _QR_PADDING + (slot_w - qr_size) // 2
    offset_y = top + _QR_PADDING + (slot_h - qr_size) // 2
    card.paste(qr_img, (offset_x, offset_y), qr_img)

    buf = io.BytesIO()
    card.save(buf, format="PNG")
    buf.seek(0)
    return buf


def safe_filename(name: str) -> str:
    safe = re.sub(r'[^\w\s-]', '', name.strip())
    safe = re.sub(r'[\s]+', '_', safe)
    return safe or "item"


def _styled_workbook(title: str, headers: list[str]) -> tuple[Workbook, any]:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    return wb, ws


def _wb_to_bytes(wb: Workbook) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ─── Startup / Seed ────────────────────────────────────────────────────────────

import time

@app.on_event("startup")
def startup() -> None:
    retries = 5
    while retries > 0:
        try:
            init_database()
            break
        except Exception as e:
            retries -= 1
            print(f"Database connection failed, retrying... ({retries} left)")
            time.sleep(3)
            
    with SessionLocal() as db:
        if settings.initial_admin_email and settings.initial_admin_password:
            email = settings.initial_admin_email.lower().strip()
            if not db.scalar(select(User).where(User.email == email)):
                db.add(User(name="Administrator", email=email, password_hash=hash_password(settings.initial_admin_password), role="superadmin"))
        elif not db.scalar(select(User.id).limit(1)):
            raise RuntimeError("No users exist. Set INITIAL_ADMIN_EMAIL and INITIAL_ADMIN_PASSWORD before first production startup.")
        db.commit()

# ─── System ─────────────────────────────────────────────────────────────────────

@app.get("/api/health", tags=["System"])
def health_check() -> dict:
    return {"status": "ok" if check_database() else "degraded", "database": check_database()}

@app.get("/api", tags=["System"])
def api_root() -> dict:
    response = {"name": "Gatherly — Event Registration & Check-in API", "version": "2.0.0"}
    if not settings.is_production:
        response["docs"] = "/docs"
    return response

# ─── Auth ───────────────────────────────────────────────────────────────────────

@app.post("/api/auth/login", tags=["Auth"])
def login(payload: LoginInput, db: Session = Depends(get_db)) -> dict:
    user = db.scalar(select(User).where(User.email == payload.email.lower().strip()))
    if not user or not user.is_active or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Email atau password tidak valid")
    return {"access_token": create_token(user.id, user.role), "token_type": "bearer", "user": UserOut.model_validate(user)}

@app.get("/api/auth/me", response_model=UserOut, tags=["Auth"])
def me(user: User = Depends(current_user)) -> User:
    return user

# ─── User Management ───────────────────────────────────────────────────────────

@app.get("/api/users", tags=["Users"])
def list_users(db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> list[dict]:
    users = db.scalars(select(User).order_by(User.name)).all()
    return [{"id": u.id, "name": u.name, "email": u.email, "role": u.role, "is_active": u.is_active, "created_at": u.created_at} for u in users]

@app.post("/api/users", tags=["Users"])
def create_user(payload: UserInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    if db.scalar(select(User).where(User.email == payload.email.lower().strip())):
        raise HTTPException(400, "Email sudah terdaftar")
    if payload.role == "superadmin" and user.role != "superadmin":
        raise HTTPException(403, "Hanya superadmin yang bisa membuat superadmin")
    new_user = User(name=payload.name, email=payload.email.lower().strip(), password_hash=hash_password(payload.password), role=payload.role)
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    return {"id": new_user.id, "name": new_user.name, "email": new_user.email, "role": new_user.role, "is_active": new_user.is_active}

@app.put("/api/users/{user_id}", tags=["Users"])
def update_user(user_id: int, payload: UserUpdateInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "User tidak ditemukan")
    if target.role == "superadmin" and user.role != "superadmin":
        raise HTTPException(403, "Hanya superadmin yang bisa mengubah superadmin")
    if payload.email.lower().strip() != target.email:
        if db.scalar(select(User).where(User.email == payload.email.lower().strip())):
            raise HTTPException(400, "Email sudah terdaftar")
    target.name = payload.name
    target.email = payload.email.lower().strip()
    target.role = payload.role
    target.is_active = payload.is_active
    db.commit()
    db.refresh(target)
    return {"id": target.id, "name": target.name, "email": target.email, "role": target.role, "is_active": target.is_active}

@app.delete("/api/users/{user_id}", tags=["Users"])
def deactivate_user(user_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "User tidak ditemukan")
    if target.id == user.id:
        raise HTTPException(400, "Tidak bisa menonaktifkan diri sendiri")
    if target.role == "superadmin" and user.role != "superadmin":
        raise HTTPException(403, "Hanya superadmin yang bisa menonaktifkan superadmin")
    target.is_active = False
    db.commit()
    return {"detail": "User berhasil dinonaktifkan"}

@app.put("/api/users/{user_id}/reset-password", tags=["Users"])
def reset_user_password(user_id: int, payload: dict, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "User tidak ditemukan")
    new_pw = payload.get("new_password", "")
    if len(new_pw) < 6:
        raise HTTPException(400, "Password minimal 6 karakter")
    target.password_hash = hash_password(new_pw)
    db.commit()
    return {"detail": "Password berhasil direset"}

@app.put("/api/auth/password", tags=["Auth"])
def change_password(payload: PasswordChangeInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(400, "Password lama tidak valid")
    user.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"detail": "Password berhasil diubah"}

# ─── Events CRUD ────────────────────────────────────────────────────────────────

@app.get("/api/events", tags=["Events"])
def list_events(db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    return [serialize_event(e, db) for e in db.scalars(select(Event).order_by(Event.event_date.desc(), Event.id.desc())).all()]

@app.post("/api/events", tags=["Events"])
def create_event(payload: EventInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    event = Event(**payload.model_dump())
    db.add(event)
    db.commit()
    db.refresh(event)
    return serialize_event(event, db)

@app.get("/api/events/{event_id}", tags=["Events"])
def get_event(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    return serialize_event(event, db)

@app.put("/api/events/{event_id}", tags=["Events"])
def update_event(event_id: int, payload: EventInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    for key, value in payload.model_dump().items():
        setattr(event, key, value)
    db.commit()
    db.refresh(event)
    return serialize_event(event, db)

@app.delete("/api/events/{event_id}", tags=["Events"])
def delete_event(event_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    db.delete(event)
    db.commit()
    return {"detail": "Event berhasil dihapus"}

@app.post("/api/events/{event_id}/copy", tags=["Events"])
def copy_event(event_id: int, payload: EventCopyInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    source = db.get(Event, event_id)
    if not source:
        raise HTTPException(404, "Event tidak ditemukan")
    new_event = Event(name=payload.name, description=source.description, event_date=payload.event_date, start_time=source.start_time, end_time=source.end_time, location=source.location, status=payload.status)
    db.add(new_event)
    db.flush()
    vendor_map = {}
    if payload.copy_vendors:
        for v in db.scalars(select(Vendor).where(Vendor.event_id == event_id)).all():
            new_v = Vendor(event_id=new_event.id, company_name=v.company_name, category=v.category, contact_name=v.contact_name, phone=v.phone, email=v.email)
            db.add(new_v)
            db.flush()
            vendor_map[v.id] = new_v.id
    table_map = {}
    if payload.copy_tables:
        for t in db.scalars(select(Table).where(Table.event_id == event_id)).all():
            new_t = Table(event_id=new_event.id, table_number=t.table_number, table_label=t.table_label, capacity=t.capacity, zone=t.zone)
            db.add(new_t)
            db.flush()
            table_map[t.id] = new_t.id
    if payload.copy_participants:
        group_token_map: dict[str, str] = {}
        for p in db.scalars(select(Participant).where(Participant.event_id == event_id)).all():
            new_group_token = None
            if p.qr_group_token:
                if p.qr_group_token not in group_token_map:
                    group_token_map[p.qr_group_token] = token_urlsafe(24)
                new_group_token = group_token_map[p.qr_group_token]
            db.add(Participant(
                event_id=new_event.id,
                vendor_id=vendor_map.get(p.vendor_id) if p.vendor_id else None,
                table_id=table_map.get(p.table_id) if p.table_id else None,
                seat_number=p.seat_number,
                name=p.name, position=p.position, phone=p.phone, email=p.email,
                qr_token=token_urlsafe(24),
                qr_group_token=new_group_token,
            ))
    db.commit()
    db.refresh(new_event)
    return serialize_event(new_event, db)

@app.get("/api/events/{event_id}/dashboard", tags=["Events"])
def dashboard(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    total = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event_id)) or 0
    checked = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event_id, Participant.check_in_at.is_not(None))) or 0
    vendors_data = []
    for vendor in db.scalars(select(Vendor).where(Vendor.event_id == event_id).order_by(Vendor.company_name)).all():
        vt = db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == vendor.id)) or 0
        vc = db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == vendor.id, Participant.check_in_at.is_not(None))) or 0
        vendors_data.append({"id": vendor.id, "company_name": vendor.company_name, "category": vendor.category, "total": vt, "checked_in": vc, "rate": round(vc / vt * 100) if vt else 0})
    recent = db.scalars(select(Participant).where(Participant.event_id == event_id, Participant.check_in_at.is_not(None)).order_by(Participant.check_in_at.desc()).limit(8)).all()
    return {"event": serialize_event(event, db), "summary": {"total": total, "checked_in": checked, "remaining": total - checked, "rate": round(checked / total * 100) if total else 0}, "vendors": vendors_data, "recent": [serialize_participant(p, db) for p in recent]}

@app.get("/api/events/{event_id}/stats", tags=["Events"])
def event_stats(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    total = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event_id)) or 0
    checked = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event_id, Participant.check_in_at.is_not(None))) or 0
    vendor_count = db.scalar(select(func.count(Vendor.id)).where(Vendor.event_id == event_id)) or 0
    table_count = db.scalar(select(func.count(Table.id)).where(Table.event_id == event_id)) or 0
    seated = db.scalar(select(func.count(Participant.id)).where(Participant.event_id == event_id, Participant.table_id.is_not(None))) or 0
    return {"event": serialize_event(event, db), "total_participants": total, "checked_in": checked, "remaining": total - checked, "rate": round(checked / total * 100) if total else 0, "vendor_count": vendor_count, "table_count": table_count, "seated_participants": seated, "unseated_participants": total - seated}

# ─── Vendors CRUD ───────────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/vendors", tags=["Vendors"])
def list_vendors(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    return [{"id": v.id, "company_name": v.company_name, "category": v.category, "contact_name": v.contact_name, "phone": v.phone, "email": v.email, "participant_count": db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == v.id)) or 0} for v in db.scalars(select(Vendor).where(Vendor.event_id == event_id).order_by(Vendor.company_name)).all()]

@app.post("/api/events/{event_id}/vendors", tags=["Vendors"])
def create_vendor(event_id: int, payload: VendorInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    vendor = Vendor(event_id=event_id, **payload.model_dump())
    db.add(vendor)
    db.commit()
    db.refresh(vendor)
    return {"id": vendor.id, **payload.model_dump(), "participant_count": 0}

@app.put("/api/events/{event_id}/vendors/{vendor_id}", tags=["Vendors"])
def update_vendor(event_id: int, vendor_id: int, payload: VendorInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    vendor = db.scalar(select(Vendor).where(Vendor.id == vendor_id, Vendor.event_id == event_id))
    if not vendor:
        raise HTTPException(404, "Vendor tidak ditemukan")
    for key, value in payload.model_dump().items():
        setattr(vendor, key, value)
    db.commit()
    db.refresh(vendor)
    return {"id": vendor.id, **payload.model_dump(), "participant_count": db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == vendor.id)) or 0}

@app.delete("/api/events/{event_id}/vendors/{vendor_id}", tags=["Vendors"])
def delete_vendor(event_id: int, vendor_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    vendor = db.scalar(select(Vendor).where(Vendor.id == vendor_id, Vendor.event_id == event_id))
    if not vendor:
        raise HTTPException(404, "Vendor tidak ditemukan")
    db.delete(vendor)
    db.commit()
    return {"detail": "Vendor berhasil dihapus"}

# ─── Tables (Seating) CRUD ─────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/tables", tags=["Seating"])
def list_tables(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    result = []
    for t in db.scalars(select(Table).where(Table.event_id == event_id).order_by(Table.table_number)).all():
        seated = db.scalar(select(func.count(Participant.id)).where(Participant.table_id == t.id)) or 0
        result.append({"id": t.id, "table_number": t.table_number, "table_label": t.table_label, "capacity": t.capacity, "zone": t.zone, "status": t.status, "seated": seated})
    return result

@app.post("/api/events/{event_id}/tables", tags=["Seating"])
def create_table(event_id: int, payload: TableInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    t = Table(event_id=event_id, **payload.model_dump())
    db.add(t)
    db.commit()
    db.refresh(t)
    return {"id": t.id, **payload.model_dump(), "seated": 0}

@app.post("/api/events/{event_id}/tables/bulk", tags=["Seating"])
def create_tables_bulk(event_id: int, payload: TableBulkInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    tables = []
    for i in range(1, payload.count + 1):
        t = Table(event_id=event_id, table_number=f"{payload.prefix}{i}", table_label=f"{payload.zone} {payload.prefix}{i}", capacity=payload.capacity, zone=payload.zone)
        db.add(t)
        tables.append(t)
    db.commit()
    return [{"id": t.id, "table_number": t.table_number, "table_label": t.table_label, "capacity": t.capacity, "zone": t.zone, "status": t.status, "seated": 0} for t in tables]

@app.put("/api/events/{event_id}/tables/{table_id}", tags=["Seating"])
def update_table(event_id: int, table_id: int, payload: TableInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    t = db.scalar(select(Table).where(Table.id == table_id, Table.event_id == event_id))
    if not t:
        raise HTTPException(404, "Meja tidak ditemukan")
    for key, value in payload.model_dump().items():
        setattr(t, key, value)
    db.commit()
    db.refresh(t)
    seated = db.scalar(select(func.count(Participant.id)).where(Participant.table_id == t.id)) or 0
    return {"id": t.id, "table_number": t.table_number, "table_label": t.table_label, "capacity": t.capacity, "zone": t.zone, "status": t.status, "seated": seated}

@app.delete("/api/events/{event_id}/tables/{table_id}", tags=["Seating"])
def delete_table(event_id: int, table_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    t = db.scalar(select(Table).where(Table.id == table_id, Table.event_id == event_id))
    if not t:
        raise HTTPException(404, "Meja tidak ditemukan")
    db.execute(select(Participant).where(Participant.table_id == table_id))
    for p in db.scalars(select(Participant).where(Participant.table_id == table_id)).all():
        p.table_id = None
        p.seat_number = None
    db.delete(t)
    db.commit()
    return {"detail": "Meja berhasil dihapus"}

# ─── Participants CRUD ──────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/participants", tags=["Participants"])
def list_participants(event_id: int, search: str = Query(""), db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    query = select(Participant).where(Participant.event_id == event_id).order_by(Participant.name)
    if search.strip():
        term = f"%{search.strip()}%"
        query = (
            query
            .outerjoin(Vendor, Participant.vendor_id == Vendor.id)
            .where(or_(
                Participant.name.ilike(term),
                Participant.email.ilike(term),
                Participant.phone.ilike(term),
                Participant.position.ilike(term),
                Vendor.company_name.ilike(term),
            ))
        )
    items = list(db.scalars(query.limit(500)).all())
    # Load full event participants once so group membership is accurate even when search filters rows
    all_for_groups = list(db.scalars(select(Participant).where(Participant.event_id == event_id)).all()) if any(p.qr_group_token for p in items) else items
    group_lookup = build_group_lookup(all_for_groups)
    return [serialize_participant(item, db, group_lookup) for item in items]

@app.post("/api/events/{event_id}/participants", tags=["Participants"])
def create_participant(event_id: int, payload: ParticipantInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    if payload.vendor_id and not db.scalar(select(Vendor).where(Vendor.id == payload.vendor_id, Vendor.event_id == event_id)):
        raise HTTPException(400, "Vendor tidak sesuai dengan event")
    if payload.table_id and not db.scalar(select(Table).where(Table.id == payload.table_id, Table.event_id == event_id)):
        raise HTTPException(400, "Meja tidak sesuai dengan event")
    participant = Participant(event_id=event_id, qr_token=token_urlsafe(24), **payload.model_dump())
    db.add(participant)
    db.commit()
    db.refresh(participant)
    return serialize_participant(participant, db)


@app.post("/api/events/{event_id}/participants/group", tags=["Participants"])
def create_participant_group(event_id: int, payload: GroupParticipantInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    _validate_group_refs(event_id, payload.vendor_id, payload.table_id, db)
    names = _normalize_group_names(payload.names)
    shared_token = token_urlsafe(24)
    created = []
    for name in names:
        p = Participant(
            event_id=event_id,
            name=name,
            vendor_id=payload.vendor_id,
            table_id=payload.table_id,
            seat_number=payload.seat_number,
            position=payload.position,
            phone=payload.phone,
            email=payload.email,
            qr_token=token_urlsafe(24),
            qr_group_token=shared_token,
        )
        db.add(p)
        created.append(p)
    db.commit()
    for p in created:
        db.refresh(p)
    return _serialize_group(created, shared_token, db)

def _normalize_group_names(raw_names: list[str]) -> list[str]:
    names = [n.strip() for n in raw_names if n and n.strip()]
    if len(names) < 2:
        raise HTTPException(400, "Delegasi minimal 2 nama")
    if any(len(n) < 2 for n in names):
        raise HTTPException(400, "Setiap nama minimal 2 karakter")
    return names


def _validate_group_refs(event_id: int, vendor_id: int | None, table_id: int | None, db: Session) -> None:
    if vendor_id and not db.scalar(select(Vendor).where(Vendor.id == vendor_id, Vendor.event_id == event_id)):
        raise HTTPException(400, "Vendor tidak sesuai dengan event")
    if table_id and not db.scalar(select(Table).where(Table.id == table_id, Table.event_id == event_id)):
        raise HTTPException(400, "Meja tidak sesuai dengan event")


def _cleanup_orphan_group_token(event_id: int, group_token: str | None, db: Session) -> None:
    if not group_token:
        return
    remaining = list(db.scalars(
        select(Participant).where(Participant.event_id == event_id, Participant.qr_group_token == group_token)
    ).all())
    if len(remaining) == 1:
        remaining[0].qr_group_token = None


def _serialize_group(members: list[Participant], shared_token: str, db: Session) -> dict:
    group_lookup = {shared_token: members} if shared_token else {}
    return {
        "qr_group_token": shared_token,
        "group_size": len(members),
        "participants": [serialize_participant(p, db, group_lookup) for p in members],
    }


@app.put("/api/events/{event_id}/participants/{participant_id}", tags=["Participants"])
def update_participant(event_id: int, participant_id: int, payload: ParticipantInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    p = db.scalar(select(Participant).where(Participant.id == participant_id, Participant.event_id == event_id))
    if not p:
        raise HTTPException(404, "Peserta tidak ditemukan")
    _validate_group_refs(event_id, payload.vendor_id, payload.table_id, db)
    for key, value in payload.model_dump().items():
        setattr(p, key, value)
    db.commit()
    db.refresh(p)
    return serialize_participant(p, db)


@app.put("/api/events/{event_id}/participants/group/{group_token}", tags=["Participants"])
def update_participant_group(event_id: int, group_token: str, payload: GroupUpdateInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    _validate_group_refs(event_id, payload.vendor_id, payload.table_id, db)
    names = _normalize_group_names(payload.names)
    members = list(db.scalars(
        select(Participant).where(
            Participant.event_id == event_id,
            Participant.qr_group_token == group_token,
        ).order_by(Participant.id).with_for_update()
    ).all())
    if not members:
        raise HTTPException(404, "Delegasi tidak ditemukan")

    shared = {
        "vendor_id": payload.vendor_id,
        "table_id": payload.table_id,
        "seat_number": payload.seat_number,
        "position": payload.position,
        "phone": payload.phone,
        "email": payload.email,
    }
    # Update overlapping rows in place (keeps attendance history)
    for idx, name in enumerate(names):
        if idx < len(members):
            members[idx].name = name
            for key, value in shared.items():
                setattr(members[idx], key, value)
        else:
            p = Participant(
                event_id=event_id,
                name=name,
                qr_token=token_urlsafe(24),
                qr_group_token=group_token,
                **shared,
            )
            db.add(p)
            members.append(p)
    # Remove excess members if names list shrank
    for extra in members[len(names):]:
        db.delete(extra)
    db.commit()

    refreshed = list(db.scalars(
        select(Participant).where(
            Participant.event_id == event_id,
            Participant.qr_group_token == group_token,
        ).order_by(Participant.id)
    ).all())
    return _serialize_group(refreshed, group_token, db)


@app.delete("/api/events/{event_id}/participants/{participant_id}", tags=["Participants"])
def delete_participant(event_id: int, participant_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    p = db.scalar(select(Participant).where(Participant.id == participant_id, Participant.event_id == event_id))
    if not p:
        raise HTTPException(404, "Peserta tidak ditemukan")
    group_token = p.qr_group_token
    db.delete(p)
    db.flush()
    _cleanup_orphan_group_token(event_id, group_token, db)
    db.commit()
    return {"detail": "Peserta berhasil dihapus"}


@app.delete("/api/events/{event_id}/participants/group/{group_token}", tags=["Participants"])
def delete_participant_group(event_id: int, group_token: str, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    members = list(db.scalars(
        select(Participant).where(
            Participant.event_id == event_id,
            Participant.qr_group_token == group_token,
        )
    ).all())
    if not members:
        raise HTTPException(404, "Delegasi tidak ditemukan")
    count = len(members)
    for member in members:
        db.delete(member)
    db.commit()
    return {"detail": "Delegasi berhasil dihapus", "deleted": count}

@app.patch("/api/events/{event_id}/participants/{participant_id}/seat", tags=["Seating"])
def assign_seat(event_id: int, participant_id: int, payload: SeatAssignInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    p = db.scalar(select(Participant).where(Participant.id == participant_id, Participant.event_id == event_id))
    if not p:
        raise HTTPException(404, "Peserta tidak ditemukan")
    if payload.table_id and not db.scalar(select(Table).where(Table.id == payload.table_id, Table.event_id == event_id)):
        raise HTTPException(400, "Meja tidak sesuai dengan event")
    p.table_id = payload.table_id
    p.seat_number = payload.seat_number
    db.commit()
    db.refresh(p)
    return serialize_participant(p, db)

# ─── Attendance ─────────────────────────────────────────────────────────────────

def _resolve_attendance_participant(event_id: int, token: str | None, participant_id: int | None, db: Session) -> Participant | None:
    if participant_id is not None:
        return db.scalar(select(Participant).where(Participant.id == participant_id, Participant.event_id == event_id))
    if token and token.strip():
        t = token.strip()
        return db.scalar(
            select(Participant).where(
                Participant.event_id == event_id,
                or_(Participant.qr_token == t, Participant.qr_group_token == t),
            )
        )
    return None


def build_checkin_preview(event_id: int, participant: Participant, db: Session) -> dict:
    members = group_members(participant, db) if participant.qr_group_token else [participant]
    is_group = bool(participant.qr_group_token) and len(members) > 1
    group_lookup = build_group_lookup(members) if is_group else None
    checked = [m for m in members if m.check_in_at]
    return {
        "result": "preview",
        "is_group": is_group,
        "group_size": len(members),
        "checked_in_count": len(checked),
        "already_checked_in_count": len(checked),
        "all_checked_in": len(checked) == len(members) and len(members) > 0,
        "participant": serialize_participant(participant, db, group_lookup),
        "participants": [serialize_participant(m, db, group_lookup) for m in members],
        "group_names": [m.name for m in members],
    }


def perform_checkin(
    event_id: int,
    participant: Participant,
    method: str,
    user: User,
    notes: str,
    db: Session,
    attended_by: str | None = None,
    is_substitute: bool = False,
    participant_ids: list[int] | None = None,
    attended_by_by_participant: dict[int, str] | None = None,
) -> dict:
    if participant.event_id != event_id:
        raise HTTPException(400, "Peserta bukan bagian dari event ini")
    if participant.qr_group_token:
        members = list(db.scalars(
            select(Participant).where(
                Participant.event_id == event_id,
                Participant.qr_group_token == participant.qr_group_token,
            ).order_by(Participant.id).with_for_update()
        ).all())
    else:
        members = [participant]

    all_members = members
    existing_checked = [member for member in all_members if member.check_in_at]
    if participant_ids is not None and len(members) > 1:
        selected = set(participant_ids)
        invalid = selected - {member.id for member in members}
        if invalid:
            raise HTTPException(400, "Anggota delegasi tidak valid")
        members = [member for member in members if member.id in selected]
        if not members:
            raise HTTPException(400, "Pilih minimal satu anggota delegasi")

    wakil_by_member = {
        int(member_id): (name or "").strip()
        for member_id, name in (attended_by_by_participant or {}).items()
    }
    invalid_wakil_ids = set(wakil_by_member) - {member.id for member in members}
    if invalid_wakil_ids:
        raise HTTPException(400, "Data wakil tidak sesuai dengan anggota yang dipilih")

    wakil_name = (attended_by or "").strip() if is_substitute else ""
    if is_substitute and not wakil_name and not wakil_by_member:
        raise HTTPException(400, "Nama wakil wajib diisi")
    if is_substitute or wakil_by_member:
        wakil_name = wakil_name or next(iter(wakil_by_member.values()), "")
    else:
        wakil_name = None

    now = datetime.now(timezone.utc)
    newly_checked: list[Participant] = []
    already: list[Participant] = []
    for member in members:
        member_wakil = wakil_by_member.get(member.id, wakil_name if is_substitute and not wakil_by_member else None)
        member_notes = notes.strip()
        if member_wakil:
            member_notes = f"Wakil: {member_wakil} atas nama {member.name}" + (f". {member_notes}" if member_notes else "")
        if member.check_in_at:
            already.append(member)
            db.add(AttendanceLog(event_id=event_id, participant_id=member.id, action="check_in", method=method, result="already_checked_in", scanned_by=user.id, notes=member_notes))
            continue
        member.attendance_status = "checked_in"
        member.check_in_at = now
        member.check_in_method = method
        member.checked_in_by = user.id
        member.attended_by = member_wakil
        newly_checked.append(member)
        db.add(AttendanceLog(event_id=event_id, participant_id=member.id, action="check_in", method=method, result="checked_in", scanned_by=user.id, notes=member_notes))
    db.commit()
    for member in members:
        db.refresh(member)
    is_group = bool(participant.qr_group_token) and len(all_members) > 1
    full_checked_count = sum(1 for member in all_members if member.check_in_at)
    pending_count = len(all_members) - full_checked_count
    if newly_checked and existing_checked:
        result = "partial_checked_in"
    elif newly_checked:
        result = "checked_in"
    else:
        result = "already_checked_in"
    primary = newly_checked[0] if newly_checked else members[0]
    return {
        "result": result,
        "is_group": is_group,
        "is_substitute": bool(wakil_by_member) or bool(wakil_name),
        "attended_by": wakil_name if not wakil_by_member or len(set(wakil_by_member.values())) == 1 else None,
        "attended_by_by_participant": {str(member.id): member.attended_by for member in all_members if member.attended_by},
        "group_size": len(all_members),
        "checked_in_count": len(newly_checked),
        "already_checked_in_count": len(existing_checked),
        "total_checked_in_count": full_checked_count,
        "remaining_count": pending_count,
        "participant": serialize_participant(primary, db),
        "participants": [serialize_participant(m, db) for m in all_members],
        "group_names": [m.name for m in all_members],
        "checked_in_at": primary.check_in_at,
    }

@app.post("/api/events/{event_id}/attendance/preview", tags=["Attendance"])
def preview_checkin(event_id: int, payload: PreviewScanInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    event = db.get(Event, event_id)
    if not event or event.status != "active":
        raise HTTPException(400, "Event tidak aktif")
    if not payload.token and payload.participant_id is None:
        raise HTTPException(400, "token atau participant_id wajib")
    participant = _resolve_attendance_participant(event_id, payload.token, payload.participant_id, db)
    if not participant:
        if payload.token:
            db.add(AttendanceLog(event_id=event_id, action="scan", method="qr", result="invalid_token", scanned_by=user.id, notes=""))
            db.commit()
        raise HTTPException(404, "QR Code tidak valid atau peserta tidak ditemukan")
    return build_checkin_preview(event_id, participant, db)

@app.post("/api/events/{event_id}/attendance/scan", tags=["Attendance"])
def scan(event_id: int, payload: ScanInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    event = db.get(Event, event_id)
    if not event or event.status != "active":
        raise HTTPException(400, "Event tidak aktif")
    token = payload.token.strip()
    participant = db.scalar(
        select(Participant).where(
            Participant.event_id == event_id,
            or_(Participant.qr_token == token, Participant.qr_group_token == token),
        ).with_for_update()
    )
    if not participant:
        db.add(AttendanceLog(event_id=event_id, action="scan", method="qr", result="invalid_token", scanned_by=user.id, notes=payload.notes))
        db.commit()
        raise HTTPException(404, "QR Code tidak valid atau peserta tidak ditemukan")
    return perform_checkin(
        event_id, participant, "qr", user, payload.notes, db,
        attended_by=payload.attended_by, is_substitute=payload.is_substitute,
        participant_ids=payload.participant_ids,
        attended_by_by_participant=payload.attended_by_by_participant,
    )

@app.post("/api/events/{event_id}/attendance/manual", tags=["Attendance"])
def manual_checkin(event_id: int, payload: ManualCheckinInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    participant = db.scalar(select(Participant).where(Participant.id == payload.participant_id, Participant.event_id == event_id).with_for_update())
    if not participant:
        raise HTTPException(404, "Peserta tidak ditemukan")
    return perform_checkin(
        event_id, participant, "manual", user, payload.notes, db,
        attended_by=payload.attended_by, is_substitute=payload.is_substitute,
        participant_ids=payload.participant_ids,
        attended_by_by_participant=payload.attended_by_by_participant,
    )

@app.post("/api/events/{event_id}/attendance/undo", tags=["Attendance"])
def undo_checkin(event_id: int, payload: UndoCheckinInput, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    p = db.scalar(select(Participant).where(Participant.id == payload.participant_id, Participant.event_id == event_id).with_for_update())
    if not p:
        raise HTTPException(404, "Peserta tidak ditemukan")
    if payload.undo_delegasi and p.qr_group_token:
        members = list(db.scalars(
            select(Participant).where(
                Participant.event_id == event_id,
                Participant.qr_group_token == p.qr_group_token,
            ).order_by(Participant.id).with_for_update()
        ).all())
    else:
        members = [p]
    undone = []
    for member in members:
        if not member.check_in_at:
            continue
        member.attendance_status = "not_attended"
        member.check_in_at = None
        member.check_in_method = None
        member.checked_in_by = None
        member.attended_by = None
        db.add(AttendanceLog(event_id=event_id, participant_id=member.id, action="undo_check_in", method="manual", result="undo", scanned_by=user.id, notes=payload.notes))
        undone.append(member)
    if not undone:
        raise HTTPException(400, "Peserta belum check-in")
    db.commit()
    for member in undone:
        db.refresh(member)
    return {
        "result": "undo_success",
        "is_group": len(members) > 1,
        "undone_count": len(undone),
        "participant": serialize_participant(undone[0], db),
        "participants": [serialize_participant(m, db) for m in members],
    }

# ─── Attendance Logs ────────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/attendance-logs", tags=["Attendance"])
def list_attendance_logs(event_id: int, limit: int = Query(100, le=500), db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict]:
    logs = db.scalars(select(AttendanceLog).where(AttendanceLog.event_id == event_id).order_by(AttendanceLog.scanned_at.desc()).limit(limit)).all()
    result = []
    for log in logs:
        p = db.get(Participant, log.participant_id) if log.participant_id else None
        scanner = db.get(User, log.scanned_by) if log.scanned_by else None
        result.append({"id": log.id, "action": log.action, "method": log.method, "result": log.result, "scanned_at": log.scanned_at, "notes": log.notes, "participant_name": p.name if p else None, "scanned_by_name": scanner.name if scanner else None})
    return result

# ─── QR Code Generation ────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/participants/qr-all", tags=["QR Code"])
def download_all_qr(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    participants_list = db.scalars(select(Participant).where(Participant.event_id == event_id).order_by(Participant.name)).all()
    if not participants_list:
        raise HTTPException(404, "Belum ada peserta di event ini")
    zip_buffer = io.BytesIO()
    seen_group_tokens: set[str] = set()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        for p in participants_list:
            members = group_members(p, db) if p.qr_group_token else [p]
            is_real_group = bool(p.qr_group_token) and len(members) > 1
            if is_real_group:
                if p.qr_group_token in seen_group_tokens:
                    continue
                seen_group_tokens.add(p.qr_group_token)
                label = "_".join(safe_filename(m.name) for m in members[:3])
                if len(members) > 3:
                    label += f"_plus{len(members) - 3}"
                token = p.qr_group_token
                filename_suffix = f"delegasi_{label}"
            else:
                token = p.qr_token
                filename_suffix = f"{safe_filename(p.name)}_{p.id}"
            img_buffer = generate_qr_image(token)
            vendor = db.get(Vendor, p.vendor_id) if p.vendor_id else None
            company = safe_filename(vendor.company_name) if vendor else "Unassigned"
            filename = f"{company}/{filename_suffix}.png"
            zf.writestr(filename, img_buffer.read())
    zip_buffer.seek(0)
    return StreamingResponse(zip_buffer, media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="qr_codes_{safe_filename(event.name)}.zip"'})

@app.get("/api/events/{event_id}/participants/{participant_id}/qr", tags=["QR Code"])
def get_participant_qr(event_id: int, participant_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    participant = db.scalar(select(Participant).where(Participant.id == participant_id, Participant.event_id == event_id))
    if not participant:
        raise HTTPException(404, "Peserta tidak ditemukan di event ini")
    buffer = generate_qr_image(effective_qr_token(participant))
    members = group_members(participant, db)
    if participant.qr_group_token and len(members) > 1:
        label = "_".join(safe_filename(m.name) for m in members[:2])
        filename = f"qr_delegasi_{label}.png"
    else:
        filename = f"qr_{safe_filename(participant.name)}_{participant.id}.png"
    return StreamingResponse(buffer, media_type="image/png", headers={"Content-Disposition": f'inline; filename="{filename}"'})

# ─── Import Excel ───────────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/import/template", tags=["Import/Export"])
def download_import_template(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")

    headers = ["nama", "perusahaan", "jabatan", "telepon", "email", "kategori_perusahaan", "nomor_meja", "nomor_kursi", "delegasi"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Template Import"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
    solo_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")       # hijau = 1 orang 1 QR
    group_fill_a = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")    # biru = delegasi A
    group_fill_b = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")    # oranye = delegasi B

    # Header WAJIB di baris 1 agar file langsung bisa di-import
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Contoh data: 2 individu + 2 grup delegasi (2 & 3 orang)
    sample_rows = [
        # hijau = solo (delegasi kosong → 1 QR sendiri)
        (["Nadia Prameswari", "PT Karya Nusantara", "Manager", "08123456789", "nadia@karya.com", "Strategic Partner", "A1", "1", ""], solo_fill),
        (["Rizky Aditya", "PT Orbit Supply", "Staff", "08198765432", "rizky@orbit.com", "Supply Chain", "A1", "2", ""], solo_fill),
        # biru = delegasi KN-VIP (2 orang, 1 QR)
        (["Edwin Sugianto", "PT Karya Nusantara", "Director", "08111111111", "edwin@karya.com", "Strategic Partner", "B1", "1", "KN-VIP"], group_fill_a),
        (["I Dewa Putu Sidan Bayupati", "PT Karya Nusantara", "Partner", "08222222222", "bayu@karya.com", "Strategic Partner", "B1", "2", "KN-VIP"], group_fill_a),
        # oranye = delegasi ORBIT-DIR (3 orang, 1 QR)
        (["Siti Rahmawati", "PT Orbit Supply", "Direktur", "08333333333", "siti@orbit.com", "Supply Chain", "C1", "1", "ORBIT-DIR"], group_fill_b),
        (["Budi Santoso", "PT Orbit Supply", "GM Operasi", "08444444444", "budi@orbit.com", "Supply Chain", "C1", "2", "ORBIT-DIR"], group_fill_b),
        (["Ayu Lestari", "PT Orbit Supply", "Sekretaris", "08555555555", "ayu@orbit.com", "Supply Chain", "C1", "3", "ORBIT-DIR"], group_fill_b),
    ]
    for r_idx, (values, fill) in enumerate(sample_rows, start=2):
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            if c_idx == 9 and val:
                cell.font = Font(bold=True, color="1565C0" if fill == group_fill_a else "E65100")

    widths = {"A": 28, "B": 22, "C": 14, "D": 14, "E": 22, "F": 18, "G": 12, "H": 12, "I": 14}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{1 + len(sample_rows)}"

    # ── Sheet Petunjuk ────────────────────────────────────────────────────────
    guide = wb.create_sheet("Petunjuk", 1)
    for col, h in enumerate(["Kolom", "Wajib", "Contoh", "Keterangan"], 1):
        cell = guide.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
    guide_rows = [
        ("nama", "Ya", "Nadia Prameswari", "Nama lengkap peserta"),
        ("perusahaan", "Ya", "PT Karya Nusantara", "Nama vendor/perusahaan (dibuat otomatis jika belum ada)"),
        ("jabatan", "Tidak", "Manager", "Posisi / jabatan"),
        ("telepon", "Tidak", "08123456789", "Nomor telepon"),
        ("email", "Tidak", "nadia@karya.com", "Email peserta"),
        ("kategori_perusahaan", "Tidak", "Strategic Partner", "Kategori vendor (default: General)"),
        ("nomor_meja", "Tidak", "A1", "Nomor meja (dibuat otomatis jika belum ada)"),
        ("nomor_kursi", "Tidak", "1", "Nomor kursi di meja"),
        ("delegasi", "Tidak", "KN-VIP", "Kode SAMA untuk anggota yang share 1 QR. Minimal 2 baris. Kosongkan jika 1 orang = 1 QR."),
    ]
    for r, row in enumerate(guide_rows, 2):
        for c, val in enumerate(row, 1):
            guide.cell(row=r, column=c, value=val)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 22
    guide.column_dimensions["D"].width = 70

    # ── Sheet Contoh Penggunaan ───────────────────────────────────────────────
    contoh = wb.create_sheet("Contoh Penggunaan", 2)
    contoh.merge_cells("A1:D1")
    judul = contoh["A1"]
    judul.value = "Gambaran hasil setelah import — berdasarkan contoh di sheet Template Import"
    judul.font = Font(bold=True, size=13, color="102A43")
    judul.alignment = Alignment(vertical="center")
    contoh.row_dimensions[1].height = 28

    for col, h in enumerate(["Jenis", "Kode Delegasi", "Anggota", "Hasil di sistem"], 1):
        cell = contoh.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    scenarios = [
        ("Individu", "(kosong)", "Nadia Prameswari", "1 undangan · 1 QR sendiri · scan absen Nadia saja", solo_fill),
        ("Individu", "(kosong)", "Rizky Aditya", "1 undangan · 1 QR sendiri · scan absen Rizky saja", solo_fill),
        ("Delegasi (2 orang)", "KN-VIP", "Edwin Sugianto + I Dewa Putu Sidan Bayupati", "1 undangan · 1 QR bersama · 1x scan = absen keduanya", group_fill_a),
        ("Delegasi (3 orang)", "ORBIT-DIR", "Siti Rahmawati + Budi Santoso + Ayu Lestari", "1 undangan · 1 QR bersama · 1x scan = absen ketiganya", group_fill_b),
    ]
    for r, (jenis, kode, anggota, hasil, fill) in enumerate(scenarios, 4):
        for c, val in enumerate([jenis, kode, anggota, hasil], 1):
            cell = contoh.cell(row=r, column=c, value=val)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        contoh.row_dimensions[r].height = 32

    contoh.cell(row=9, column=1, value="Cara pakai singkat:").font = Font(bold=True, size=11)
    steps = [
        "1. Download template ini, buka sheet Template Import.",
        "2. Hapus baris contoh (baris hijau/biru/oranye), atau timpa dengan data asli.",
        "3. Untuk peserta sendiri: biarkan kolom delegasi KOSONG.",
        "4. Untuk delegasi: isi kode yang SAMA di kolom delegasi (contoh KN-VIP) pada semua anggota grup (min. 2 nama).",
        "5. Simpan sebagai .xlsx → di aplikasi klik Import Excel → cek preview → Confirm.",
        "6. Download all QR: tiap individu 1 file; tiap kode delegasi juga 1 file QR saja.",
    ]
    for i, step in enumerate(steps, 10):
        contoh.cell(row=i, column=1, value=step)
        contoh.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)

    contoh.column_dimensions["A"].width = 22
    contoh.column_dimensions["B"].width = 16
    contoh.column_dimensions["C"].width = 48
    contoh.column_dimensions["D"].width = 55

    buf = _wb_to_bytes(wb)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template_import_peserta.xlsx"'},
    )

@app.post("/api/events/{event_id}/import/preview", tags=["Import/Export"])
async def import_preview(event_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not db.get(Event, event_id):
        raise HTTPException(404, "Event tidak ditemukan")
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "File harus berformat .xlsx")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception:
        raise HTTPException(400, "File Excel tidak valid")
    headers_row = [str(cell.value or "").strip().lower() for cell in ws[1]]
    required = {"nama", "perusahaan"}
    if not required.issubset(set(headers_row)):
        raise HTTPException(400, f"Header wajib tidak ditemukan: {required - set(headers_row)}")
    col_map = {h: i for i, h in enumerate(headers_row)}
    rows = []
    errors = []
    existing_vendors = {v.company_name.lower() for v in db.scalars(select(Vendor).where(Vendor.event_id == event_id)).all()}
    existing_tables = {t.table_number.lower() for t in db.scalars(select(Table).where(Table.event_id == event_id)).all()}
    new_vendors = set()
    new_tables = set()
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[col_map.get("nama", 0)]:
            continue
        name = str(row[col_map["nama"]]).strip()
        company = str(row[col_map["perusahaan"]]).strip() if col_map.get("perusahaan") is not None and row[col_map["perusahaan"]] else ""
        if not name:
            errors.append({"row": row_idx, "error": "Nama kosong"})
            continue
        if not company:
            errors.append({"row": row_idx, "error": "Perusahaan kosong"})
            continue
        row_data = {
            "nama": name,
            "perusahaan": company,
            "jabatan": str(row[col_map.get("jabatan", -1)] or "").strip() if col_map.get("jabatan") is not None and len(row) > col_map.get("jabatan", 999) else "",
            "telepon": str(row[col_map.get("telepon", -1)] or "").strip() if col_map.get("telepon") is not None and len(row) > col_map.get("telepon", 999) else "",
            "email": str(row[col_map.get("email", -1)] or "").strip() if col_map.get("email") is not None and len(row) > col_map.get("email", 999) else "",
            "kategori_perusahaan": str(row[col_map.get("kategori_perusahaan", -1)] or "General").strip() if col_map.get("kategori_perusahaan") is not None and len(row) > col_map.get("kategori_perusahaan", 999) else "General",
            "nomor_meja": str(row[col_map.get("nomor_meja", -1)] or "").strip() if col_map.get("nomor_meja") is not None and len(row) > col_map.get("nomor_meja", 999) else "",
            "nomor_kursi": str(row[col_map.get("nomor_kursi", -1)] or "").strip() if col_map.get("nomor_kursi") is not None and len(row) > col_map.get("nomor_kursi", 999) else "",
            "delegasi": (
                str(row[col_map["delegasi"]] or "").strip()
                if col_map.get("delegasi") is not None and len(row) > col_map["delegasi"]
                else (
                    str(row[col_map["grup"]] or "").strip()
                    if col_map.get("grup") is not None and len(row) > col_map["grup"]
                    else ""
                )
            ),
        }
        # Check for duplicate in DB
        existing = db.scalar(select(Participant).where(Participant.event_id == event_id, func.lower(Participant.name) == name.lower()).join(Vendor, Participant.vendor_id == Vendor.id, isouter=True).where(or_(Vendor.company_name.ilike(company), Participant.vendor_id.is_(None))))
        row_data["is_duplicate"] = existing is not None
        if company.lower() not in existing_vendors:
            new_vendors.add(company)
        if row_data["nomor_meja"] and row_data["nomor_meja"].lower() not in existing_tables:
            new_tables.add(row_data["nomor_meja"])
        rows.append(row_data)

    # Validate delegasi groups: same code must appear on ≥2 rows
    delegasi_counts = Counter((r.get("delegasi") or "").strip() for r in rows if (r.get("delegasi") or "").strip())
    delegasi_groups = []
    for key, count in sorted(delegasi_counts.items()):
        names = [r["nama"] for r in rows if (r.get("delegasi") or "").strip() == key]
        delegasi_groups.append({"kode": key, "jumlah": count, "nama": names})
        if count < 2:
            errors.append({"row": "-", "error": f"Delegasi '{key}' hanya {count} orang — minimal 2 nama dengan kode yang sama (akan diabaikan, jadi QR individual)"})

    return {
        "total_rows": len(rows),
        "errors": errors,
        "rows": rows,
        "new_vendors": sorted(new_vendors),
        "new_tables": sorted(new_tables),
        "delegasi_groups": delegasi_groups,
    }

@app.post("/api/events/{event_id}/import/confirm", tags=["Import/Export"])
def import_confirm(event_id: int, payload: ImportConfirmInput, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))) -> dict:
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    created = 0
    skipped = 0
    updated = 0
    vendors_created = 0
    tables_created = 0
    import_group_tokens: dict[str, str] = {}
    for row in payload.rows:
        company = row.get("perusahaan", "")
        vendor = db.scalar(select(Vendor).where(Vendor.event_id == event_id, func.lower(Vendor.company_name) == company.lower())) if company else None
        if company and not vendor:
            vendor = Vendor(event_id=event_id, company_name=company, category=row.get("kategori_perusahaan", "General"))
            db.add(vendor)
            db.flush()
            vendors_created += 1
        table = None
        table_number = row.get("nomor_meja", "")
        if table_number:
            table = db.scalar(select(Table).where(Table.event_id == event_id, func.lower(Table.table_number) == table_number.lower()))
            if not table:
                table = Table(event_id=event_id, table_number=table_number, capacity=8, zone="Regular")
                db.add(table)
                db.flush()
                tables_created += 1
        name = row.get("nama", "")
        group_key = (row.get("delegasi") or row.get("grup") or "").strip()
        group_token = None
        if group_key:
            if group_key not in import_group_tokens:
                import_group_tokens[group_key] = token_urlsafe(24)
            group_token = import_group_tokens[group_key]
        if row.get("is_duplicate") and payload.duplicate_action == "skip":
            skipped += 1
            continue
        if row.get("is_duplicate") and payload.duplicate_action == "update":
            existing = db.scalar(select(Participant).where(Participant.event_id == event_id, func.lower(Participant.name) == name.lower()))
            if existing:
                existing.position = row.get("jabatan", existing.position)
                existing.phone = row.get("telepon", existing.phone)
                existing.email = row.get("email", existing.email)
                existing.vendor_id = vendor.id if vendor else existing.vendor_id
                existing.table_id = table.id if table else existing.table_id
                existing.seat_number = row.get("nomor_kursi") or existing.seat_number
                if group_token:
                    existing.qr_group_token = group_token
                updated += 1
                continue
        p = Participant(
            event_id=event_id,
            vendor_id=vendor.id if vendor else None,
            table_id=table.id if table else None,
            seat_number=row.get("nomor_kursi") or None,
            name=name,
            position=row.get("jabatan", ""),
            phone=row.get("telepon", ""),
            email=row.get("email", ""),
            qr_token=token_urlsafe(24),
            qr_group_token=group_token,
        )
        db.add(p)
        created += 1

    # Drop group tokens that ended up with fewer than 2 members (not a real delegasi)
    db.flush()
    for group_token in import_group_tokens.values():
        members = list(db.scalars(
            select(Participant).where(Participant.event_id == event_id, Participant.qr_group_token == group_token)
        ).all())
        if len(members) < 2:
            for m in members:
                m.qr_group_token = None

    db.commit()
    return {"created": created, "skipped": skipped, "updated": updated, "vendors_created": vendors_created, "tables_created": tables_created}

# ─── Export Excel ───────────────────────────────────────────────────────────────

@app.get("/api/events/{event_id}/export/participants", tags=["Import/Export"])
def export_participants(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    headers = ["No", "Nama", "Perusahaan", "Jabatan", "Telepon", "Email", "Meja", "Kursi", "Delegasi", "Anggota Delegasi", "Status", "Hadir Sebagai", "Waktu Check-in"]
    wb, ws = _styled_workbook("Daftar Peserta", headers)
    people = list(db.scalars(select(Participant).where(Participant.event_id == event_id).order_by(Participant.name)).all())
    group_lookup = build_group_lookup(people)
    # Stable human labels (DEL-1, DEL-2, ...) so export can be re-imported
    delegasi_labels: dict[str, str] = {}
    for p in people:
        if not p.qr_group_token or p.qr_group_token in delegasi_labels:
            continue
        mates = group_lookup.get(p.qr_group_token, [p])
        if len(mates) > 1:
            delegasi_labels[p.qr_group_token] = f"DEL-{len(delegasi_labels) + 1}"
    for idx, p in enumerate(people, 1):
        vendor = db.get(Vendor, p.vendor_id) if p.vendor_id else None
        table = db.get(Table, p.table_id) if p.table_id else None
        delegasi_label = delegasi_labels.get(p.qr_group_token or "", "")
        anggota = ""
        if delegasi_label:
            mates = group_lookup.get(p.qr_group_token, [p])
            anggota = ", ".join(m.name for m in mates)
        ws.append([
            idx, p.name, vendor.company_name if vendor else "", p.position, p.phone, p.email,
            table.table_number if table else "", p.seat_number or "",
            delegasi_label, anggota,
            "Hadir" if p.check_in_at else "Belum Hadir",
            p.attended_by or "",
            str(p.check_in_at)[:19] if p.check_in_at else "",
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    ws.column_dimensions["J"].width = 40
    buf = _wb_to_bytes(wb)
    fn = f"peserta_{safe_filename(event.name)}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}"'})

@app.get("/api/events/{event_id}/export/attendance", tags=["Import/Export"])
def export_attendance(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    headers = ["No", "Nama", "Perusahaan", "Delegasi", "Hadir Sebagai", "Waktu Check-in", "Metode", "Meja", "Kursi", "Dicatat Oleh"]
    wb, ws = _styled_workbook("Laporan Kehadiran", headers)
    all_people = list(db.scalars(select(Participant).where(Participant.event_id == event_id)).all())
    group_lookup = build_group_lookup(all_people)
    delegasi_labels: dict[str, str] = {}
    for p in all_people:
        if not p.qr_group_token or p.qr_group_token in delegasi_labels:
            continue
        mates = group_lookup.get(p.qr_group_token, [p])
        if len(mates) > 1:
            delegasi_labels[p.qr_group_token] = f"DEL-{len(delegasi_labels) + 1}"
    checked = db.scalars(select(Participant).where(Participant.event_id == event_id, Participant.check_in_at.is_not(None)).order_by(Participant.check_in_at)).all()
    for idx, p in enumerate(checked, 1):
        vendor = db.get(Vendor, p.vendor_id) if p.vendor_id else None
        table = db.get(Table, p.table_id) if p.table_id else None
        scanner = db.get(User, p.checked_in_by) if p.checked_in_by else None
        ws.append([
            idx, p.name, vendor.company_name if vendor else "",
            delegasi_labels.get(p.qr_group_token or "", ""),
            p.attended_by or "",
            str(p.check_in_at)[:19] if p.check_in_at else "", p.check_in_method or "",
            table.table_number if table else "", p.seat_number or "", scanner.name if scanner else "",
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = _wb_to_bytes(wb)
    fn = f"kehadiran_{safe_filename(event.name)}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}"'})

@app.get("/api/events/{event_id}/export/vendors", tags=["Import/Export"])
def export_vendors(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    headers = ["No", "Perusahaan", "Kategori", "Kontak", "Telepon", "Email", "Jumlah Peserta", "Hadir", "Rate (%)"]
    wb, ws = _styled_workbook("Daftar Vendor", headers)
    for idx, v in enumerate(db.scalars(select(Vendor).where(Vendor.event_id == event_id).order_by(Vendor.company_name)).all(), 1):
        total = db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == v.id)) or 0
        checked = db.scalar(select(func.count(Participant.id)).where(Participant.vendor_id == v.id, Participant.check_in_at.is_not(None))) or 0
        ws.append([idx, v.company_name, v.category, v.contact_name, v.phone, v.email, total, checked, round(checked / total * 100) if total else 0])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = _wb_to_bytes(wb)
    fn = f"vendor_{safe_filename(event.name)}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}"'})

@app.get("/api/events/{event_id}/export/seating", tags=["Import/Export"])
def export_seating(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    headers = ["Meja", "Label", "Zona", "Kapasitas", "Terisi", "Peserta", "Kursi"]
    wb, ws = _styled_workbook("Layout Meja", headers)
    for t in db.scalars(select(Table).where(Table.event_id == event_id).order_by(Table.table_number)).all():
        seated_list = db.scalars(select(Participant).where(Participant.table_id == t.id).order_by(Participant.seat_number)).all()
        if not seated_list:
            ws.append([t.table_number, t.table_label, t.zone, t.capacity, 0, "", ""])
        for p in seated_list:
            ws.append([t.table_number, t.table_label, t.zone, t.capacity, len(seated_list), p.name, p.seat_number or ""])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = _wb_to_bytes(wb)
    fn = f"seating_{safe_filename(event.name)}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}"'})

@app.get("/api/events/{event_id}/export/audit-log", tags=["Import/Export"])
def export_audit_log(event_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    event = db.get(Event, event_id)
    if not event:
        raise HTTPException(404, "Event tidak ditemukan")
    headers = ["No", "Waktu", "Aksi", "Metode", "Hasil", "Peserta", "Dicatat Oleh", "Catatan"]
    wb, ws = _styled_workbook("Audit Log", headers)
    for idx, log in enumerate(db.scalars(select(AttendanceLog).where(AttendanceLog.event_id == event_id).order_by(AttendanceLog.scanned_at.desc())).all(), 1):
        p = db.get(Participant, log.participant_id) if log.participant_id else None
        scanner = db.get(User, log.scanned_by) if log.scanned_by else None
        ws.append([idx, str(log.scanned_at)[:19] if log.scanned_at else "", log.action, log.method, log.result, p.name if p else "-", scanner.name if scanner else "-", log.notes])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = _wb_to_bytes(wb)
    fn = f"audit_{safe_filename(event.name)}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{fn}"'})
